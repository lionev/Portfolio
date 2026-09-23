# -*- coding: utf-8 -*-
"""Generate Julio Jorge cover letter in PT/EN as DOCX and PDF."""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "cv"
DOWNLOADS = Path.home() / "Downloads" / "Julio_Jorge" / "cartas"
FONTS = Path(r"C:\Windows\Fonts")
SITE = "https://juliojorge.dev"

LETTER = {
    "pt": {
        "kicker": "Carta de apresentação",
        "role": "Desenvolvedor Full Stack",
        "date": "Recife, 22 de setembro de 2026",
        "greeting": "Prezados,",
        "paragraphs": [
            "Sou Julio Jorge, desenvolvedor full stack em Recife. Estou aberto a vagas de produto, sistemas internos e ferramentas desktop — o tipo de trabalho em que o código precisa chegar na mão de quem usa.",
            "No IASSEPE fui o único desenvolvedor da intranet institucional — front-end, back-end, banco de dados e implantação: SPAs com React 19, TypeScript e Tailwind, e APIs REST com Node.js, Express, JWT, LDAP e PostgreSQL com Prisma. Também desenvolvi e implantei um sistema em Google Apps Script para gerenciar mudanças em planilhas e arquivos no Google Sheets e no Google Drive. Antes disso, como gerente de redes, cuidei de disponibilidade, segurança e da coordenação da equipe de TI.",
            "Em paralelo construo produtos próprios. O NAVO é um marketplace de agenda para tatuagem, barbeiro e estética, com busca por proximidade e sinal no PIX. O XoHub e o XoMacro são ferramentas desktop e web que uma comunidade de jogos abre no dia a dia.",
            f"O portfólio com esses projetos, a trajetória e o currículo em português e inglês está no ar em {SITE}. Vale abrir o site: ali dá para ver o que eu entrego, sem depender só desta carta.",
            "Gosto de fechar o ciclo — do requisito ao que aparece na tela, e do que quebra em produção quando não tem ninguém olhando. Fico à disposição para conversar.",
        ],
        "closing": "Atenciosamente,",
        "sign": "Julio Jorge",
        "docx": "Julio_Jorge_Carta.docx",
        "pdf": "Julio_Jorge_Carta.pdf",
        "title": "Julio Jorge — Carta de apresentação",
    },
    "en": {
        "kicker": "Cover letter",
        "role": "Full Stack Developer",
        "date": "Recife, 22 September 2026",
        "greeting": "Dear hiring team,",
        "paragraphs": [
            "I am Julio Jorge, a full stack developer based in Recife, Brazil. I am open to product, internal systems and desktop-tool roles — work where the code has to reach the person who actually uses it.",
            "At IASSEPE I was the sole developer of the institutional intranet — front-end, back-end, database and deployment: SPAs with React 19, TypeScript and Tailwind, and REST APIs with Node.js, Express, JWT, LDAP and PostgreSQL with Prisma. I also designed and deployed a Google Apps Script system to manage changes in spreadsheets and files on Google Sheets and Google Drive. Before that, as network manager, I owned availability, security and coordination of the IT team.",
            "In parallel I build my own products. NAVO is a booking marketplace for tattoo, barber and aesthetics shops, with nearby search and a PIX deposit to lock a slot. XoHub and XoMacro are desktop and web tools that a gaming community opens every day.",
            f"The portfolio with that work, my path and the résumé in Portuguese and English is live at {SITE}. The site is the shortest way to see what I ship, beyond this letter.",
            "I like to close the loop — from the requirement to what shows on screen, and from what breaks in production when nobody is watching. I would welcome a conversation.",
        ],
        "closing": "Kind regards,",
        "sign": "Julio Jorge",
        "docx": "Julio_Jorge_Cover_Letter.docx",
        "pdf": "Julio_Jorge_Cover_Letter.pdf",
        "title": "Julio Jorge — Cover letter",
    },
}


def rgb(hex_color: str) -> RGBColor:
    hex_color = hex_color.lstrip("#")
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def set_run_font(run, name="Calibri"):
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:cs"), name)
    rFonts.set(qn("w:eastAsia"), name)


def add_run(p, text, *, size, color, bold=False):
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    set_run_font(run)
    return run


def add_hyperlink(paragraph, text, url, *, size, color):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")
    rFonts.set(qn("w:cs"), "Calibri")
    rFonts.set(qn("w:eastAsia"), "Calibri")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), color.lstrip("#"))
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.extend([rFonts, sz, color_el, u])
    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def build_docx(lang: str, dest: Path) -> Path:
    data = LETTER[lang]
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Twips(40)
    add_run(p, "Julio Jorge", size=20, color="1A1A1A", bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Twips(20)
    add_run(p, data["role"], size=11, color="555555")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Twips(12)
    add_run(p, "julionev10@gmail.com  ·  +55 81 98627-3135  ·  Recife, PE", size=10, color="666666")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Twips(80)
    add_hyperlink(p, SITE.replace("https://", ""), SITE, size=10, color="1A1A1A")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Twips(120)
    p.paragraph_format.space_after = Twips(200)
    add_run(p, data["kicker"].upper(), size=9, color="888888", bold=True)
    add_run(p, "   ·   ", size=9, color="BBBBBB")
    add_run(p, data["date"], size=9, color="888888")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Twips(160)
    add_run(p, data["greeting"], size=11, color="1A1A1A")

    for text in data["paragraphs"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Twips(160)
        p.paragraph_format.line_spacing = 1.18
        if SITE in text:
            before, after = text.split(SITE, 1)
            add_run(p, before, size=11, color="333333")
            add_hyperlink(p, SITE, SITE, size=11, color="1A1A1A")
            add_run(p, after, size=11, color="333333")
        else:
            add_run(p, text, size=11, color="333333")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Twips(80)
    p.paragraph_format.space_after = Twips(40)
    add_run(p, data["closing"], size=11, color="333333")

    p = doc.add_paragraph()
    add_run(p, data["sign"], size=12, color="1A1A1A", bold=True)

    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(dest)
    return dest


def register_fonts():
    pdfmetrics.registerFont(TTFont("Calibri", str(FONTS / "calibri.ttf")))
    pdfmetrics.registerFont(TTFont("Calibri-Bold", str(FONTS / "calibrib.ttf")))


def build_pdf(lang: str, dest: Path) -> Path:
    data = LETTER[lang]
    dest.parent.mkdir(parents=True, exist_ok=True)

    styles = {
        "name": ParagraphStyle(
            "name",
            fontName="Calibri-Bold",
            fontSize=20,
            textColor=HexColor("#1A1A1A"),
            leading=24,
            spaceAfter=2,
        ),
        "role": ParagraphStyle(
            "role",
            fontName="Calibri",
            fontSize=11,
            textColor=HexColor("#555555"),
            leading=14,
            spaceAfter=4,
        ),
        "meta": ParagraphStyle(
            "meta",
            fontName="Calibri",
            fontSize=10,
            textColor=HexColor("#666666"),
            leading=13,
            spaceAfter=2,
        ),
        "link": ParagraphStyle(
            "link",
            fontName="Calibri",
            fontSize=10,
            textColor=HexColor("#1A1A1A"),
            leading=13,
            spaceAfter=16,
        ),
        "kicker": ParagraphStyle(
            "kicker",
            fontName="Calibri",
            fontSize=9,
            textColor=HexColor("#888888"),
            leading=12,
            spaceBefore=8,
            spaceAfter=16,
        ),
        "body": ParagraphStyle(
            "body",
            fontName="Calibri",
            fontSize=11,
            textColor=HexColor("#333333"),
            leading=16,
            spaceAfter=11,
        ),
        "close": ParagraphStyle(
            "close",
            fontName="Calibri",
            fontSize=11,
            textColor=HexColor("#333333"),
            leading=15,
            spaceBefore=6,
            spaceAfter=4,
        ),
        "sign": ParagraphStyle(
            "sign",
            fontName="Calibri-Bold",
            fontSize=12,
            textColor=HexColor("#1A1A1A"),
            leading=15,
        ),
    }

    story = [
        Paragraph("Julio Jorge", styles["name"]),
        Paragraph(data["role"], styles["role"]),
        Paragraph("julionev10@gmail.com  ·  +55 81 98627-3135  ·  Recife, PE", styles["meta"]),
        Paragraph(f'<link href="{SITE}"><u>juliojorge.dev</u></link>', styles["link"]),
        Paragraph(f'{data["kicker"].upper()}   ·   {data["date"]}', styles["kicker"]),
        Paragraph(data["greeting"], styles["body"]),
    ]
    for text in data["paragraphs"]:
        html = text.replace(SITE, f'<link href="{SITE}"><u>{SITE}</u></link>')
        story.append(Paragraph(html, styles["body"]))
    story.append(Paragraph(data["closing"], styles["close"]))
    story.append(Paragraph(data["sign"], styles["sign"]))

    doc = SimpleDocTemplate(
        str(dest),
        pagesize=letter,
        leftMargin=1.0 * inch,
        rightMargin=1.0 * inch,
        topMargin=0.9 * inch,
        bottomMargin=0.9 * inch,
        title=data["title"],
        author="Julio Jorge",
    )
    doc.build(story)
    return dest


def main():
    register_fonts()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    outputs = []
    for lang in ("pt", "en"):
        data = LETTER[lang]
        outputs.append(build_docx(lang, OUT_DIR / data["docx"]))
        outputs.append(build_pdf(lang, OUT_DIR / data["pdf"]))

    DOWNLOADS.mkdir(parents=True, exist_ok=True)
    copies = []
    for path in outputs:
        dest = DOWNLOADS / path.name
        shutil.copy2(path, dest)
        copies.append(dest)

    print("Generated:")
    for path in outputs:
        print(" ", path, path.stat().st_size, "bytes")
    print("Downloads:")
    for path in copies:
        print(" ", path)


if __name__ == "__main__":
    main()
