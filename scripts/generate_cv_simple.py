# -*- coding: utf-8 -*-
"""Simple one-column Julio Jorge résumé (PT/EN) — DOCX + PDF."""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_JUSTIFY, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "cv-simple"
PUBLIC = ROOT / "public"
DOWNLOADS = Path.home() / "Downloads" / "Julio_Jorge" / "modelo-simples"
FONTS = Path(r"C:\Windows\Fonts")

INK = "1A1A1A"
MUTED = "555555"
SOFT = "666666"
INK_HEX = HexColor("#1A1A1A")
MUTED_HEX = HexColor("#555555")
SOFT_HEX = HexColor("#666666")


def jobs_pt() -> list[dict]:
    return [
        {
            "title": "Desenvolvedor Full Stack",
            "dates": "mar/2026 – jul/2026",
            "place": "IASSEPE — Recife, PE",
            "bullets": [
                "Desenvolvimento individual da intranet institucional — front-end, back-end, banco de dados e implantação: SPAs com React 19, TypeScript, Vite e Tailwind CSS; APIs REST com Node.js, Express, JWT, LDAP e PostgreSQL + Prisma.",
                "Sistema em Google Apps Script, também desenvolvido e implantado por mim, para gerenciar mudanças em planilhas e arquivos no Google Sheets e no Google Drive.",
            ],
            "skills": "React 19 · TypeScript · Vite · Tailwind CSS · Node.js · Express · JWT · LDAP · PostgreSQL · Prisma · Google Apps Script · Google Sheets · Google Drive",
        },
        {
            "title": "Gerente de Redes",
            "dates": "jun/2025 – mar/2026",
            "place": "IASSEPE — Recife, PE",
            "bullets": [
                "Gestão de redes corporativas com foco em disponibilidade e desempenho.",
                "Liderança da equipe de TI em projetos de expansão e modernização.",
                "Administração de TCP/IP, BGP, OSPF, MPLS, SD-WAN, firewalls, IDS/IPS e ZTNA.",
            ],
        },
        {
            "title": "Desenvolvedor Full Stack",
            "dates": "2026",
            "place": "NAVO — Produto próprio, Recife, PE",
            "bullets": [
                "Marketplace de agenda com sinal no PIX para tatuagem, barbeiro e estética.",
                "Painel do dono, pagamentos via Mercado Pago e busca por proximidade.",
                "Desenvolvimento auxiliado por IA generativa na escrita e revisão do código.",
            ],
            "skills": "Next.js 16 · React 19 · TypeScript · Prisma · PostgreSQL · Mercado Pago · Zustand · IA generativa",
        },
        {
            "title": "Desenvolvedor Front-End",
            "dates": "2026",
            "place": "Portfólio — juliojorge.dev",
            "bullets": [
                "Site pessoal bilíngue em React 19, TypeScript, Vite, Tailwind CSS e Motion.",
                "Download do currículo em PDF com escolha de idioma independente da página.",
            ],
            "skills": "React 19 · TypeScript · Vite · Tailwind CSS · Motion · Vercel",
        },
        {
            "title": "Desenvolvedor Full Stack",
            "dates": "2025–2026",
            "place": "XoHub — Produto próprio, Recife, PE",
            "bullets": [
                "Hub de módulos para Path of Exile: HUD desktop em Python/PySide6 e painel web em React.",
                "API Node.js/Express com Prisma e PostgreSQL; módulos por assinatura e atualização automática.",
                "Desenvolvimento auxiliado por IA generativa na escrita e revisão do código.",
            ],
            "skills": "Python 3 · PySide6 · React · TypeScript · Vite · Tailwind CSS · Node.js · Express · Prisma · PostgreSQL · IA generativa",
        },
        {
            "title": "Desenvolvedor Desktop",
            "dates": "2026",
            "place": "XoMacro — Produto próprio, Recife, PE",
            "bullets": [
                "Gravador e repetidor de teclado e mouse para Windows em C# / .NET 8.",
                "Hotkeys globais, temas claro/escuro e publicação autocontida até a build 62.",
                "Desenvolvimento auxiliado por IA generativa na escrita e revisão do código.",
            ],
            "skills": "C# · .NET 8 · WinForms · Hotkeys · Simulação de input · IA generativa",
        },
        {
            "title": "Desenvolvedor Full Stack",
            "dates": "2024–2026",
            "place": "XoD (PKMDBOT) — Produto próprio, Recife, PE",
            "bullets": [
                "Site, planos, launcher e bot desktop — xod.vercel.app.",
                "Mercado Pago, indicação e automação por visão de tela no cliente Windows.",
                "Desenvolvimento auxiliado por IA generativa na escrita e revisão do código.",
            ],
            "skills": "React · TypeScript · Node.js · Express · Mercado Pago · Python 3 · OpenCV · CustomTkinter · PyInstaller · IA generativa",
        },
        {
            "title": "Desenvolvedor",
            "dates": "Produto próprio",
            "place": "XoBot Twitch Manager — Recife, PE",
            "bullets": [
                "Gerenciador de bots para Twitch com OAuth 2.0, Helix API e IRC.",
                "Back-end remoto via Google Apps Script e distribuição em executável Windows.",
                "Desenvolvimento auxiliado por IA generativa na escrita e revisão do código.",
            ],
            "skills": "Python 3 · Tkinter · asyncio · Twitch API · OAuth 2.0 · Google Apps Script · PyInstaller · IA generativa",
        },
        {
            "title": "Analista de Suporte Júnior",
            "dates": "jul/2024 – mar/2025",
            "place": "Diário de Pernambuco — Recife, PE",
            "bullets": [
                "Atendimento de suporte N1, N2 e N3, com resolução de chamados via Help Desk.",
                "Implantação de ERP comercial e industrial e suporte ao domínio Active Directory.",
            ],
        },
        {
            "title": "Analista de Suporte de TI",
            "dates": "fev/2022 – jun/2023",
            "place": "Conceitto Industrial (terceirizado) — Recife, PE",
            "bullets": [
                "Suporte N1, N2 e N3 por telefone, e-mail, tickets e acesso remoto.",
                "Implantação de ERP e identificação de falhas com parecer técnico.",
            ],
        },
        {
            "title": "Estagiário de TI",
            "dates": "mar/2020 – fev/2022",
            "place": "Grande Recife Consórcio de Transporte — Recife, PE",
            "bullets": [
                "Manutenção de computadores desktop, suporte à rede e ao domínio corporativo.",
                "Instalação e suporte de sistemas de terceiros.",
            ],
        },
        {
            "title": "Estagiário de TI",
            "dates": "mar/2018 – fev/2020",
            "place": "Prefeitura da Cidade do Paulista — Paulista, PE",
            "bullets": [
                "Manutenção de computadores no ambiente público municipal.",
                "Instalação e suporte de sistemas para os departamentos da prefeitura.",
            ],
        },
    ]


def jobs_en() -> list[dict]:
    return [
        {
            "title": "Full Stack Developer",
            "dates": "Mar/2026 – Jul/2026",
            "place": "IASSEPE — Recife, PE, Brazil",
            "bullets": [
                "Sole developer of the institutional intranet — front-end, back-end, database and deployment: SPAs with React 19, TypeScript, Vite and Tailwind CSS; REST APIs with Node.js, Express, JWT, LDAP and PostgreSQL + Prisma.",
                "Google Apps Script system, also designed and deployed by me, to manage changes in spreadsheets and files on Google Sheets and Google Drive.",
            ],
            "skills": "React 19 · TypeScript · Vite · Tailwind CSS · Node.js · Express · JWT · LDAP · PostgreSQL · Prisma · Google Apps Script · Google Sheets · Google Drive",
        },
        {
            "title": "Network Manager",
            "dates": "Jun/2025 – Mar/2026",
            "place": "IASSEPE — Recife, PE, Brazil",
            "bullets": [
                "Corporate network infrastructure with a focus on availability and performance.",
                "Led the IT team on expansion and modernization projects.",
                "TCP/IP, BGP, OSPF, MPLS, SD-WAN, firewalls, IDS/IPS and ZTNA.",
            ],
        },
        {
            "title": "Full Stack Developer",
            "dates": "2026",
            "place": "NAVO — Independent product, Recife, PE",
            "bullets": [
                "Booking marketplace with a PIX deposit for tattoo, barber and aesthetics shops.",
                "Owner panel, Mercado Pago payments and nearby search.",
                "Generative AI used as a development aid for writing and reviewing the code.",
            ],
            "skills": "Next.js 16 · React 19 · TypeScript · Prisma · PostgreSQL · Mercado Pago · Zustand · Generative AI",
        },
        {
            "title": "Front-End Developer",
            "dates": "2026",
            "place": "Portfolio — juliojorge.dev",
            "bullets": [
                "Bilingual personal site in React 19, TypeScript, Vite, Tailwind CSS and Motion.",
                "PDF résumé download whose language does not follow the page language.",
            ],
            "skills": "React 19 · TypeScript · Vite · Tailwind CSS · Motion · Vercel",
        },
        {
            "title": "Full Stack Developer",
            "dates": "2025–2026",
            "place": "XoHub — Independent product, Recife, PE",
            "bullets": [
                "Module hub for Path of Exile: Python/PySide6 desktop HUD and a React web panel.",
                "Node.js/Express API with Prisma and PostgreSQL; subscription modules and auto-update.",
                "Generative AI used as a development aid for writing and reviewing the code.",
            ],
            "skills": "Python 3 · PySide6 · React · TypeScript · Vite · Tailwind CSS · Node.js · Express · Prisma · PostgreSQL · Generative AI",
        },
        {
            "title": "Desktop Developer",
            "dates": "2026",
            "place": "XoMacro — Independent product, Recife, PE",
            "bullets": [
                "Keyboard and mouse recorder/player for Windows in C# / .NET 8.",
                "Global hotkeys, light/dark themes and a self-contained publish through build 62.",
                "Generative AI used as a development aid for writing and reviewing the code.",
            ],
            "skills": "C# · .NET 8 · WinForms · Hotkeys · Input simulation · Generative AI",
        },
        {
            "title": "Full Stack Developer",
            "dates": "2024–2026",
            "place": "XoD (PKMDBOT) — Independent product, Recife, PE",
            "bullets": [
                "Site, plans, launcher and desktop bot — xod.vercel.app.",
                "Mercado Pago, referrals and screen-vision automation on the Windows client.",
                "Generative AI used as a development aid for writing and reviewing the code.",
            ],
            "skills": "React · TypeScript · Node.js · Express · Mercado Pago · Python 3 · OpenCV · CustomTkinter · PyInstaller · Generative AI",
        },
        {
            "title": "Developer",
            "dates": "Independent product",
            "place": "XoBot Twitch Manager — Recife, PE",
            "bullets": [
                "Twitch bot manager with OAuth 2.0, Helix API and IRC.",
                "Remote back-end via Google Apps Script, shipped as a Windows executable.",
                "Generative AI used as a development aid for writing and reviewing the code.",
            ],
            "skills": "Python 3 · Tkinter · asyncio · Twitch API · OAuth 2.0 · Google Apps Script · PyInstaller · Generative AI",
        },
        {
            "title": "Junior Support Analyst",
            "dates": "Jul/2024 – Mar/2025",
            "place": "Diário de Pernambuco — Recife, PE, Brazil",
            "bullets": [
                "N1, N2 and N3 support, resolving tickets through Help Desk.",
                "Commercial and industrial ERP rollout and Active Directory domain support.",
            ],
        },
        {
            "title": "IT Support Analyst",
            "dates": "Feb/2022 – Jun/2023",
            "place": "Conceitto Industrial (outsourced) — Recife, PE, Brazil",
            "bullets": [
                "N1, N2 and N3 support by phone, email, tickets and remote access.",
                "ERP rollout and failure identification with technical reports.",
            ],
        },
        {
            "title": "IT Intern",
            "dates": "Mar/2020 – Feb/2022",
            "place": "Grande Recife Transit Consortium — Recife, PE, Brazil",
            "bullets": [
                "Desktop maintenance, network support and the corporate domain.",
                "Installation and support of third-party systems.",
            ],
        },
        {
            "title": "IT Intern",
            "dates": "Mar/2018 – Feb/2020",
            "place": "Paulista City Hall — Paulista, PE, Brazil",
            "bullets": [
                "Computer maintenance in a municipal public environment.",
                "Installation and support of systems for city hall departments.",
            ],
        },
    ]


CV = {
    "pt": {
        "city": "Recife, PE",
        "contact": "+55 81 98627-3135  |  julionev10@gmail.com  |  juliojorge.dev",
        "profile_h": "RESUMO PROFISSIONAL",
        "profile": (
            "Desenvolvedor Full Stack e Gerente de Redes, com cerca de 8 anos em tecnologia — "
            "do suporte N1–N3 e da infraestrutura até intranet, sistemas internos e produtos no ar. "
            "Desenvolvimento de ponta a ponta com React, Next.js, TypeScript, Node.js, Express, "
            "PostgreSQL, Prisma, Python e C# / .NET, além de redes TCP/IP, BGP, OSPF, MPLS, "
            "SD-WAN e segurança (firewall, IDS/IPS, ZTNA). Experiência em APIs REST, autenticação "
            "(JWT, LDAP), Google Apps Script (Sheets e Drive), pagamentos (Mercado Pago e PIX), "
            "ferramentas desktop e IA generativa "
            "como apoio ao desenvolvimento. Perfil orientado a entrega, disponibilidade e ao que "
            "chega na mão de quem usa."
        ),
        "exp_h": "EXPERIÊNCIA PROFISSIONAL",
        "skills_label": "Stack: ",
        "edu_h": "ESCOLARIDADE",
        "edu": "Bacharelado em Sistemas da Informação  |  UNINASSAU — Recife, PE  |  2024",
        "skills_h": "HABILIDADES",
        "skills_left": [
            "JavaScript / TypeScript",
            "React / Next.js",
            "Node.js / Express",
            "PostgreSQL / Prisma",
            "Python 3 / PySide6",
            "C# / .NET 8",
            "Vite / Tailwind CSS",
            "Google Apps Script",
            "IA generativa",
        ],
        "skills_right": [
            "TCP/IP, BGP, OSPF, MPLS",
            "SD-WAN / Firewall",
            "IDS/IPS / ZTNA",
            "Active Directory",
            "Suporte N1 / N2 / N3",
            "ERP / Help Desk",
            "Virtualização",
            "Gestão de equipe",
        ],
        "courses_h": "CURSOS",
        "courses": [
            "Dev Salesforce — Apex, Trigger, Flow e LWC  |  out/2026  |  Udemy · 5 horas",
            "Salesforce na prática para administradores  |  out/2026  |  Udemy · 3,5 horas",
            "QA e Testes Automatizados  |  2024  |  Udemy — Cypress básico e intermediário",
            "Desenvolvedor Front-End  |  2024  |  Rocketseat — React, Java, NLWs",
        ],
        "jobs": jobs_pt(),
        "docx": "Julio_Jorge_Curriculo.docx",
        "pdf": "Julio_Jorge_Curriculo.pdf",
        "title": "Julio Jorge — Currículo",
    },
    "en": {
        "city": "Recife, PE — Brazil",
        "contact": "+55 81 98627-3135  |  julionev10@gmail.com  |  juliojorge.dev",
        "profile_h": "PROFESSIONAL SUMMARY",
        "profile": (
            "Full stack developer and network manager, with about 8 years in technology — "
            "from N1–N3 support and infrastructure to intranet, internal systems and live products. "
            "End-to-end development with React, Next.js, TypeScript, Node.js, Express, "
            "PostgreSQL, Prisma, Python and C# / .NET, plus TCP/IP, BGP, OSPF, MPLS, SD-WAN "
            "and security (firewall, IDS/IPS, ZTNA). Experience with REST APIs, authentication "
            "(JWT, LDAP), Google Apps Script (Sheets and Drive), payments (Mercado Pago and PIX), "
            "desktop tools and generative AI as "
            "a development aid. Oriented toward shipping, availability and what reaches the "
            "person who uses it."
        ),
        "exp_h": "PROFESSIONAL EXPERIENCE",
        "skills_label": "Stack: ",
        "edu_h": "EDUCATION",
        "edu": "B.Sc. in Information Systems  |  UNINASSAU — Recife, PE  |  2024",
        "skills_h": "SKILLS",
        "skills_left": [
            "JavaScript / TypeScript",
            "React / Next.js",
            "Node.js / Express",
            "PostgreSQL / Prisma",
            "Python 3 / PySide6",
            "C# / .NET 8",
            "Vite / Tailwind CSS",
            "Google Apps Script",
            "Generative AI",
        ],
        "skills_right": [
            "TCP/IP, BGP, OSPF, MPLS",
            "SD-WAN / Firewall",
            "IDS/IPS / ZTNA",
            "Active Directory",
            "Support N1 / N2 / N3",
            "ERP / Help Desk",
            "Virtualization",
            "Team leadership",
        ],
        "courses_h": "COURSES",
        "courses": [
            "Salesforce developer — Apex, Trigger, Flow and LWC  |  Oct 2026  |  Udemy · 5 hours",
            "Salesforce in practice for admins  |  Oct 2026  |  Udemy · 3.5 hours",
            "QA & automated testing  |  2024  |  Udemy — Cypress beginner and intermediate",
            "Front-End Developer  |  2024  |  Rocketseat — React, Java, NLWs",
        ],
        "jobs": jobs_en(),
        "docx": "Julio_Jorge_Resume.docx",
        "pdf": "Julio_Jorge_Resume.pdf",
        "title": "Julio Jorge — Resume",
    },
}


def rgb(hex_color: str) -> RGBColor:
    hex_color = hex_color.lstrip("#")
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def set_run_font(run, name="Calibri"):
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:cs"), name)
    rFonts.set(qn("w:eastAsia"), name)


def shade_none(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:color"), "FFFFFF")
        borders.append(el)
    tcPr.append(borders)


def set_cell_margins(cell, top=0, left=0, bottom=0, right=0):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for name, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{name}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        tcMar.append(el)
    tcPr.append(tcMar)


def set_table_width(table, width):
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(width))
    tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd")
        tblPr.append(tblInd)
    tblInd.set(qn("w:w"), "0")
    tblInd.set(qn("w:type"), "dxa")
    jc = tblPr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        tblPr.append(jc)
    jc.set(qn("w:val"), "left")


def add_run(p, text, *, size, color, bold=False):
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    set_run_font(run)
    return run


def style_p(p, *, before=0, after=40, align=WD_ALIGN_PARAGRAPH.LEFT):
    pf = p.paragraph_format
    pf.space_before = Twips(before)
    pf.space_after = Twips(after)
    pf.line_spacing = 1.08
    pf.alignment = align
    return p


def add_heading(doc, text, *, first=False):
    return add_row_table(doc, text, "", left_bold=True, before=80 if first else 220, after=80)


def add_body(doc, text, *, before=0, after=60):
    p = doc.add_paragraph()
    style_p(p, before=before, after=after)
    add_run(p, text, size=10.5, color=MUTED)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph()
    style_p(p, before=0, after=28)
    pf = p.paragraph_format
    pf.left_indent = Inches(0.18)
    add_run(p, "•  " + text, size=10.5, color=MUTED)
    return p


def add_skills(doc, label, text):
    p = doc.add_paragraph()
    style_p(p, before=12, after=20)
    add_run(p, label + text, size=9, color=SOFT)
    return p


def add_row_table(doc, left_text, right_text, *, left_bold=True, before=140, after=8, left_w=6552, right_w=3528):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.allow_autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_width(table, left_w + right_w)
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for child in list(grid):
            grid.remove(child)
        for w in (left_w, right_w):
            col = OxmlElement("w:gridCol")
            col.set(qn("w:w"), str(w))
            grid.append(col)
    left, right = table.rows[0].cells
    for cell, width, align, text, bold, color, size in (
        (left, left_w, WD_ALIGN_PARAGRAPH.LEFT, left_text, left_bold, INK, 11),
        (right, right_w, WD_ALIGN_PARAGRAPH.RIGHT, right_text, False, MUTED, 10.5),
    ):
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        tcW = tcPr.find(qn("w:tcW"))
        if tcW is None:
            tcW = OxmlElement("w:tcW")
            tcPr.append(tcW)
        tcW.set(qn("w:w"), str(width))
        tcW.set(qn("w:type"), "dxa")
        shade_none(cell)
        set_cell_margins(cell, 0, 0, 0, 0)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        style_p(p, before=before, after=after, align=align)
        if text:
            add_run(p, text, size=size, color=color, bold=bold)
    return table


def build_docx(lang: str, dest: Path) -> Path:
    data = CV[lang]
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.6)

    name = doc.add_paragraph()
    style_p(name, before=0, after=20)
    add_run(name, "Julio Jorge", size=26, color=INK, bold=True)

    add_row_table(doc, data["city"], data["contact"], left_bold=False, before=0, left_w=3024, right_w=7056)
    left, right = doc.tables[0].rows[0].cells
    for cell, size, color in ((left, 10.5, SOFT), (right, 10.5, SOFT)):
        run = cell.paragraphs[0].runs[0]
        run.bold = False
        run.font.size = Pt(size)
        run.font.color.rgb = rgb(color)

    add_heading(doc, data["profile_h"], first=True)
    add_body(doc, data["profile"], after=40)

    add_heading(doc, data["exp_h"])
    for i, job in enumerate(data["jobs"]):
        add_row_table(doc, job["title"], job["dates"], before=160 if i else 40)
        add_body(doc, job["place"], before=0, after=20)
        for bullet in job["bullets"]:
            add_bullet(doc, bullet)
        if job.get("skills"):
            add_skills(doc, data["skills_label"], job["skills"])

    add_heading(doc, data["edu_h"])
    add_body(doc, data["edu"], after=20)

    add_heading(doc, data["skills_h"])
    skills = doc.add_table(rows=1, cols=2)
    skills.autofit = False
    skills.allow_autofit = False
    set_table_width(skills, 10800)
    left, right = skills.rows[0].cells
    for cell, items, pad_left in (
        (left, data["skills_left"], 0),
        (right, data["skills_right"], 200),
    ):
        shade_none(cell)
        set_cell_margins(cell, 0, pad_left, 0, 80)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        first = True
        for item in items:
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            style_p(p, before=0, after=28)
            add_run(p, "•  " + item, size=10.5, color=MUTED)

    add_heading(doc, data["courses_h"])
    for course in data["courses"]:
        add_body(doc, course, after=20)

    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(dest)
    return dest


def register_fonts():
    pdfmetrics.registerFont(TTFont("Calibri", str(FONTS / "calibri.ttf")))
    pdfmetrics.registerFont(TTFont("Calibri-Bold", str(FONTS / "calibrib.ttf")))


def styles():
    return {
        "name": ParagraphStyle(
            "name",
            fontName="Calibri-Bold",
            fontSize=26,
            textColor=INK_HEX,
            leading=30,
            spaceAfter=2,
        ),
        "city": ParagraphStyle(
            "city",
            fontName="Calibri",
            fontSize=10.5,
            textColor=SOFT_HEX,
            leading=13,
            alignment=TA_LEFT,
        ),
        "contact": ParagraphStyle(
            "contact",
            fontName="Calibri",
            fontSize=10.5,
            textColor=SOFT_HEX,
            leading=13,
            alignment=TA_RIGHT,
        ),
        "h": ParagraphStyle(
            "h",
            fontName="Calibri-Bold",
            fontSize=11,
            textColor=INK_HEX,
            leading=14,
            spaceBefore=2,
            spaceAfter=2,
        ),
        "profile": ParagraphStyle(
            "profile",
            fontName="Calibri",
            fontSize=10.5,
            textColor=MUTED_HEX,
            leading=14,
            alignment=TA_JUSTIFY,
            spaceAfter=2,
        ),
        "job": ParagraphStyle(
            "job",
            fontName="Calibri-Bold",
            fontSize=11,
            textColor=INK_HEX,
            leading=14,
            alignment=TA_LEFT,
        ),
        "dates": ParagraphStyle(
            "dates",
            fontName="Calibri",
            fontSize=10.5,
            textColor=MUTED_HEX,
            leading=14,
            alignment=TA_RIGHT,
        ),
        "place": ParagraphStyle(
            "place",
            fontName="Calibri",
            fontSize=10.5,
            textColor=SOFT_HEX,
            leading=13,
            spaceAfter=1,
        ),
        "body": ParagraphStyle(
            "body",
            fontName="Calibri",
            fontSize=10.5,
            textColor=MUTED_HEX,
            leading=13.5,
            leftIndent=10,
            spaceAfter=1.2,
        ),
        "plain": ParagraphStyle(
            "plain",
            fontName="Calibri",
            fontSize=10.5,
            textColor=MUTED_HEX,
            leading=13.5,
            spaceAfter=1.5,
        ),
        "skill": ParagraphStyle(
            "skill",
            fontName="Calibri",
            fontSize=10.5,
            textColor=MUTED_HEX,
            leading=14,
            spaceAfter=1.4,
        ),
        "jobskills": ParagraphStyle(
            "jobskills",
            fontName="Calibri",
            fontSize=9,
            textColor=SOFT_HEX,
            leading=12,
            spaceBefore=1,
            spaceAfter=2,
        ),
    }


LINE_COLS = [4.55 * inch, 2.45 * inch]


def line_table(left, right, *, top=0, bottom=0, valign="MIDDLE"):
    table = Table([[left, right]], colWidths=LINE_COLS)
    table.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), valign),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                ("TOPPADDING", (0, 0), (-1, -1), top),
                ("BOTTOMPADDING", (0, 0), (-1, -1), bottom),
            ]
        )
    )
    return table


def header_table(data, s):
    table = Table(
        [[Paragraph(data["city"], s["city"]), Paragraph(data["contact"], s["contact"])]],
        colWidths=[2.1 * inch, 4.9 * inch],
    )
    table.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                ("TOPPADDING", (0, 0), (-1, -1), 0),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ]
        )
    )
    return table


def heading_table(text, s):
    return line_table(Paragraph(text, s["h"]), "", top=8, bottom=2)


def job_table(job, s):
    return line_table(
        Paragraph(job["title"], s["job"]),
        Paragraph(job["dates"], s["dates"]),
        top=8,
        valign="BOTTOM",
    )


def skills_table(data, s):
    left = [Paragraph("•  " + item, s["skill"]) for item in data["skills_left"]]
    right = [Paragraph("•  " + item, s["skill"]) for item in data["skills_right"]]
    rows = []
    for i in range(max(len(left), len(right))):
        rows.append(
            [
                left[i] if i < len(left) else "",
                right[i] if i < len(right) else "",
            ]
        )
    table = Table(rows, colWidths=[3.5 * inch, 3.5 * inch])
    table.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (0, -1), 0),
                ("LEFTPADDING", (1, 0), (1, -1), 12),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 0),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
            ]
        )
    )
    return table


def build_pdf(lang: str, dest: Path) -> Path:
    data = CV[lang]
    dest.parent.mkdir(parents=True, exist_ok=True)
    s = styles()
    story = [
        Paragraph("Julio Jorge", s["name"]),
        header_table(data, s),
        Spacer(1, 8),
        heading_table(data["profile_h"], s),
        Paragraph(data["profile"], s["profile"]),
        heading_table(data["exp_h"], s),
    ]
    for job in data["jobs"]:
        block = [
            job_table(job, s),
            Paragraph(job["place"], s["place"]),
        ]
        for bullet in job["bullets"]:
            block.append(Paragraph("•  " + bullet, s["body"]))
        if job.get("skills"):
            block.append(Paragraph(data["skills_label"] + job["skills"], s["jobskills"]))
        story.append(KeepTogether(block))

    story.append(heading_table(data["edu_h"], s))
    story.append(Paragraph(data["edu"], s["plain"]))
    story.append(heading_table(data["skills_h"], s))
    story.append(skills_table(data, s))
    story.append(heading_table(data["courses_h"], s))
    for course in data["courses"]:
        story.append(Paragraph(course, s["plain"]))

    doc = SimpleDocTemplate(
        str(dest),
        pagesize=letter,
        leftMargin=0.75 * inch,
        rightMargin=0.75 * inch,
        topMargin=0.6 * inch,
        bottomMargin=0.55 * inch,
        title=data["title"],
        author="Julio Jorge",
    )
    doc.build(story)
    return dest


def main():
    register_fonts()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    DOWNLOADS.mkdir(parents=True, exist_ok=True)
    outputs = []
    for lang in ("pt", "en"):
        data = CV[lang]
        outputs.append(build_docx(lang, OUT_DIR / data["docx"]))
        outputs.append(build_pdf(lang, OUT_DIR / data["pdf"]))

    PUBLIC.mkdir(parents=True, exist_ok=True)
    site = {
        "pt": PUBLIC / "JulioJorge.pdf",
        "en": PUBLIC / "JulioJorgeResume.pdf",
    }
    pack = {
        "pt": DOWNLOADS / "JulioJorge.pdf",
        "en": DOWNLOADS / "JulioJorgeResume.pdf",
    }
    shutil.copy2(OUT_DIR / CV["pt"]["pdf"], site["pt"])
    shutil.copy2(OUT_DIR / CV["en"]["pdf"], site["en"])

    copies = []
    for lang, dest in pack.items():
        shutil.copy2(OUT_DIR / CV[lang]["pdf"], dest)
        copies.append(dest)
        docx_src = OUT_DIR / CV[lang]["docx"]
        docx_dest = dest.with_suffix(".docx")
        shutil.copy2(docx_src, docx_dest)
        copies.append(docx_dest)

    for leftover in (
        PUBLIC / "Julio_Jorge.pdf",
        PUBLIC / "Julio_Jorge_Resume.pdf",
    ):
        leftover.unlink(missing_ok=True)

    print("Generated:")
    for path in outputs:
        print(" ", path, path.stat().st_size, "bytes")
    print("Site:")
    for path in site.values():
        print(" ", path)
    print("Downloads:")
    for path in copies:
        print(" ", path)


if __name__ == "__main__":
    main()
