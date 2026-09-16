# -*- coding: utf-8 -*-
"""Generate Julio Jorge resume in PT/EN as DOCX and PDF."""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    KeepTogether,
    PageTemplate,
    Paragraph,
    Spacer,
)

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "cv"
DOWNLOADS = Path.home() / "Downloads"
PUBLIC = ROOT / "public"
FONTS = Path(r"C:\Windows\Fonts")

SIDEBAR_W = 3100  # twips
MAIN_W = 9140
PAGE_W = 12240
PAGE_H = 15840

SEP = " · "


def jobs_pt() -> list[dict]:
    return [
        {
            "title": "Desenvolvedor Full Stack  —  IASSEPE  ·  mar/2026 – jul/2026  ·  5 meses",
            "subtitle": "Instituto de Atenção à Saúde e Bem-estar dos Servidores do Estado de Pernambuco",
            "bullets": [
                "Desenvolvimento de sistemas internos para o serviço público, de ponta a ponta: SPAs com React 19, TypeScript, Vite e Tailwind CSS no front-end; APIs REST com Node.js, Express, JWT, LDAP e PostgreSQL + Prisma no back-end.",
            ],
            "skills": "React 19 · TypeScript · Vite · Tailwind CSS · Radix UI · shadcn/ui · React Hook Form · Zod · Node.js · Express · JWT · LDAP · PostgreSQL · Prisma",
        },
        {
            "title": "Gerente de Redes  —  IASSEPE  ·  jun/2025 – mar/2026  ·  10 meses",
            "subtitle": "",
            "bullets": [
                "Gestão de infraestrutura de redes corporativas com foco em disponibilidade e desempenho.",
                "Liderança da equipe de TI, com coordenação de projetos de expansão e modernização.",
                "Implementação e administração de redes TCP/IP, BGP, OSPF, MPLS, SD-WAN e ambientes virtualizados.",
                "Configuração e manutenção de firewalls, IDS/IPS e soluções de segurança ZTNA.",
            ],
            "skills": "TCP/IP · BGP · OSPF · MPLS · SD-WAN · Firewall · IDS/IPS · ZTNA · Virtualização · Gestão de infraestrutura · Liderança de equipe",
        },
        {
            "title": "Desenvolvedor Full Stack  —  XoHub  ·  2025–2026  ·  Produto próprio",
            "subtitle": "Hub de módulos para Path of Exile 1 e 2 — overlay desktop e painel web por assinatura",
            "bullets": [
                "Concepção e evolução de um hub completo: HUD desktop em Python/PySide6 e painel web em React para liberar módulos por assinatura.",
                "Cliente com Cluster Jewels, regex de vendor/stash, PoB Trade Links, Blight, Farms com Atlas, Timeless Jewels, Fast Leveling e Auto Flask no PoE 2.",
                "Price check que lê o item na tela, consulta o trade oficial e o poe.ninja e abre a busca com filtros reais, com limite de requisições e cache.",
                "API Node.js/Express com Prisma e PostgreSQL; cada módulo pode ser adquirido isoladamente, em ciclos de 1 a 3 meses, com atualização automática via version.json.",
            ],
            "skills": "Python 3 · PySide6 · React · TypeScript · Vite · Tailwind CSS · Node.js · Express · Prisma · PostgreSQL",
        },
        {
            "title": "Desenvolvedor Desktop  —  XoMacro  ·  2026  ·  Produto próprio",
            "subtitle": "Gravador e repetidor de teclado e mouse para Windows — C# / .NET 8",
            "bullets": [
                "Aplicativo Windows com biblioteca de macros, disparo por tecla ou botão do mouse, foco em processo específico e repetição por contagem, ao manter pressionado ou em loop infinito.",
                "Passos combináveis (combo, texto, espera, clique) ou linha do tempo de eventos gravados, com retomada por estágio e velocidade de 0,25× a 3×.",
                "Hotkeys globais, watermark de status na tela, bandeja do sistema, DPI por monitor, temas claro/escuro e publicação autocontida em net8.0-windows até a build 62.",
            ],
            "skills": "C# · .NET 8 · WinForms · Hotkeys · Simulação de input",
        },
        {
            "title": "Desenvolvedor Full Stack  —  XoD (PKMDBOT)  ·  2024–2026  ·  Produto próprio",
            "subtitle": "Produto com site, planos, launcher e bot desktop — xod.vercel.app",
            "bullets": [
                "Aplicação web com seleção dinâmica de planos conforme o status do personagem (novo, teste ativo, teste expirado ou assinante).",
                "Integração com Mercado Pago (Preference API e webhook), Google Sheets como base e sistema de indicação validado no front-end e no back-end.",
                "Cliente desktop (XoD-Bot) com automação por visão de tela, Mini HUD nativo e distribuição Windows via PyInstaller.",
            ],
            "skills": "React · TypeScript · Node.js · Express · Mercado Pago · Python 3 · OpenCV · CustomTkinter · PyInstaller",
        },
        {
            "title": "Desenvolvedor  —  XoBot Twitch Manager  ·  Produto próprio",
            "subtitle": "",
            "bullets": [
                "Gerenciador de bots para Twitch com interface em Python, OAuth 2.0, Helix API e IRC.",
                "Back-end remoto via Google Apps Script e Google Sheets para onboarding e renovação automática de tokens, com distribuição em executável Windows.",
            ],
            "skills": "Python 3 · Tkinter · asyncio · Twitch API · OAuth 2.0 · Google Apps Script · PyInstaller",
        },
        {
            "title": "Analista de Suporte Júnior  —  Diário de Pernambuco  ·  jul/2024 – mar/2025  ·  9 meses",
            "subtitle": "",
            "bullets": [
                "Atendimento de suporte N1, N2 e N3, com resolução de chamados via Help Desk.",
                "Implantação de ERP comercial e industrial, adequando o sistema às necessidades do cliente.",
                "Identificação de falhas, emissão de pareceres técnicos e suporte ao domínio Active Directory.",
            ],
            "skills": "Suporte N1/N2/N3 · Help Desk · ERP · Active Directory · Windows · Análise de falhas",
        },
        {
            "title": "Analista de Suporte de TI  —  Conceitto Industrial (Terceirizado)  ·  fev/2022 – jun/2023  ·  1 ano e 5 meses",
            "subtitle": "",
            "bullets": [
                "Suporte N1, N2 e N3 por telefone, e-mail, sistema de tickets e acesso remoto.",
                "Implantação de ERP comercial e industrial e identificação de falhas com emissão de parecer técnico.",
            ],
            "skills": "Suporte N1/N2/N3 · Help Desk · ERP · Acesso remoto · Tickets · Análise de falhas",
        },
        {
            "title": "Estagiário de TI  —  Grande Recife Consórcio de Transporte  ·  mar/2020 – fev/2022  ·  2 anos",
            "subtitle": "",
            "bullets": [
                "Manutenção de computadores desktop, suporte à rede e ao domínio corporativo.",
                "Instalação e suporte de sistemas de terceiros.",
            ],
            "skills": "Manutenção de hardware · Suporte técnico · Redes · Active Directory",
        },
        {
            "title": "Estagiário de TI  —  Prefeitura da Cidade do Paulista  ·  mar/2018 – fev/2020  ·  2 anos",
            "subtitle": "",
            "bullets": [
                "Manutenção de computadores no ambiente público municipal.",
                "Instalação e suporte de sistemas para os departamentos da prefeitura.",
            ],
            "skills": "Manutenção de hardware · Suporte técnico · Instalação de sistemas",
        },
    ]


def jobs_en() -> list[dict]:
    return [
        {
            "title": "Full Stack Developer  —  IASSEPE  ·  Mar/2026 – Jul/2026  ·  5 months",
            "subtitle": "Institute for Health Care and Well-being of Civil Servants of the State of Pernambuco",
            "bullets": [
                "Internal systems for public service, end to end: SPAs with React 19, TypeScript, Vite and Tailwind CSS on the front-end; REST APIs with Node.js, Express, JWT, LDAP and PostgreSQL + Prisma on the back-end.",
            ],
            "skills": "React 19 · TypeScript · Vite · Tailwind CSS · Radix UI · shadcn/ui · React Hook Form · Zod · Node.js · Express · JWT · LDAP · PostgreSQL · Prisma",
        },
        {
            "title": "Network Manager  —  IASSEPE  ·  Jun/2025 – Mar/2026  ·  10 months",
            "subtitle": "",
            "bullets": [
                "Management of corporate network infrastructure with a focus on availability and performance.",
                "Leadership of the IT team, coordinating expansion and modernization projects.",
                "Implementation and administration of TCP/IP, BGP, OSPF, MPLS and SD-WAN networks, plus virtualized environments.",
                "Configuration and maintenance of firewalls, IDS/IPS and ZTNA security solutions.",
            ],
            "skills": "TCP/IP · BGP · OSPF · MPLS · SD-WAN · Firewall · IDS/IPS · ZTNA · Virtualization · Infrastructure management · Team leadership",
        },
        {
            "title": "Full Stack Developer  —  XoHub  ·  2025–2026  ·  Independent product",
            "subtitle": "Module hub for Path of Exile 1 and 2 — desktop overlay and subscription web panel",
            "bullets": [
                "Designed and evolved a full hub: a Python/PySide6 desktop HUD and a React web panel to unlock modules by subscription.",
                "Client modules include Cluster Jewels, vendor/stash regex, PoB Trade Links, Blight, Farms with Atlas, Timeless Jewels, Fast Leveling and Auto Flask on PoE 2.",
                "Price check reads the item on screen, queries official trade and poe.ninja, and opens search with real filters, rate limiting and cache.",
                "Node.js/Express API with Prisma and PostgreSQL; each module can be purchased on its own, on a 1 to 3 month cycle, with automatic updates via version.json.",
            ],
            "skills": "Python 3 · PySide6 · React · TypeScript · Vite · Tailwind CSS · Node.js · Express · Prisma · PostgreSQL",
        },
        {
            "title": "Desktop Developer  —  XoMacro  ·  2026  ·  Independent product",
            "subtitle": "Keyboard and mouse recorder/player for Windows — C# / .NET 8",
            "bullets": [
                "Windows app with a macro library, triggers from a key or mouse button, focus on a specific process, and repeat by count, hold or infinite loop.",
                "Each macro mixes steps (combo, text, wait, click) or a timeline of recorded events, with resume by stage and speed from 0.25× to 3×.",
                "Global hotkeys, on-screen status watermark, system tray, per-monitor DPI, light/dark themes and a self-contained net8.0-windows publish through build 62.",
            ],
            "skills": "C# · .NET 8 · WinForms · Hotkeys · Input simulation",
        },
        {
            "title": "Full Stack Developer  —  XoD (PKMDBOT)  ·  2024–2026  ·  Independent product",
            "subtitle": "Product with site, plans, launcher and desktop bot — xod.vercel.app",
            "bullets": [
                "Web app with dynamic plan selection based on character status (new, active trial, expired trial or subscriber).",
                "Mercado Pago integration (Preference API and webhook), Google Sheets as the store, and a referral system validated on both front-end and back-end.",
                "Desktop client (XoD-Bot) with screen-vision automation, a native Mini HUD and Windows distribution via PyInstaller.",
            ],
            "skills": "React · TypeScript · Node.js · Express · Mercado Pago · Python 3 · OpenCV · CustomTkinter · PyInstaller",
        },
        {
            "title": "Developer  —  XoBot Twitch Manager  ·  Independent product",
            "subtitle": "",
            "bullets": [
                "Twitch bot manager with a Python GUI, OAuth 2.0, Helix API and IRC.",
                "Remote back-end via Google Apps Script and Google Sheets for onboarding and automatic token refresh, distributed as a Windows executable.",
            ],
            "skills": "Python 3 · Tkinter · asyncio · Twitch API · OAuth 2.0 · Google Apps Script · PyInstaller",
        },
        {
            "title": "Junior Support Analyst  —  Diário de Pernambuco  ·  Jul/2024 – Mar/2025  ·  9 months",
            "subtitle": "",
            "bullets": [
                "N1, N2 and N3 support, resolving tickets through Help Desk.",
                "Rollout of commercial and industrial ERP systems tailored to the client’s needs.",
                "Failure identification, technical reports and Active Directory domain support.",
            ],
            "skills": "Support N1/N2/N3 · Help Desk · ERP · Active Directory · Windows · Failure analysis",
        },
        {
            "title": "IT Support Analyst  —  Conceitto Industrial (Outsourced)  ·  Feb/2022 – Jun/2023  ·  1 year 5 months",
            "subtitle": "",
            "bullets": [
                "N1, N2 and N3 support by phone, email, ticketing system and remote access.",
                "Commercial and industrial ERP rollout and failure identification with technical reports.",
            ],
            "skills": "Support N1/N2/N3 · Help Desk · ERP · Remote access · Tickets · Failure analysis",
        },
        {
            "title": "IT Intern  —  Grande Recife Consórcio de Transporte  ·  Mar/2020 – Feb/2022  ·  2 years",
            "subtitle": "",
            "bullets": [
                "Desktop maintenance, network support and the corporate domain.",
                "Installation and support of third-party systems.",
            ],
            "skills": "Hardware maintenance · Technical support · Networking · Active Directory",
        },
        {
            "title": "IT Intern  —  Paulista City Hall  ·  Mar/2018 – Feb/2020  ·  2 years",
            "subtitle": "",
            "bullets": [
                "Computer maintenance in a municipal public environment.",
                "Installation and support of systems for city hall departments.",
            ],
            "skills": "Hardware maintenance · Technical support · System installation",
        },
    ]


def projects_pt() -> list[dict]:
    return [
        {
            "title": "PoeRedux  —  Fork WPF / C#",
            "bullets": [
                "Fork de um patcher visual para Path of Exile: detecção automática da pasta Steam/GGG, presets, UI em português, save único e log do que cada patch alterou.",
            ],
            "skills": "C# · WPF · .NET · JSON",
        },
        {
            "title": "sCocktail  —  Front-end",
            "bullets": [
                "Catálogo de mais de 100 drinques via API, com interface enxuta em React e Tailwind.",
            ],
            "skills": "TypeScript · React · Vite · Tailwind · Axios",
        },
        {
            "title": "Timer  —  Front-end",
            "bullets": [
                "Pomodoro em React com navegação via Reducer e Context para manter o foco.",
            ],
            "skills": "TypeScript · React · Styled Components · Immer",
        },
        {
            "title": "To-Do  —  Front-end",
            "bullets": [
                "Gestão de tarefas com 9 componentes padronizados publicados via Tsup.",
            ],
            "skills": "TypeScript · React",
        },
    ]


def projects_en() -> list[dict]:
    return [
        {
            "title": "PoeRedux  —  WPF / C# fork",
            "bullets": [
                "Fork of a visual Path of Exile patcher: automatic Steam/GGG folder detection, presets, Portuguese UI, a single save and a log of what each patch changed.",
            ],
            "skills": "C# · WPF · .NET · JSON",
        },
        {
            "title": "sCocktail  —  Front-end",
            "bullets": [
                "Catalog of 100+ drinks via API, with a lean React and Tailwind interface.",
            ],
            "skills": "TypeScript · React · Vite · Tailwind · Axios",
        },
        {
            "title": "Timer  —  Front-end",
            "bullets": [
                "Pomodoro timer in React with Reducer and Context navigation to stay focused.",
            ],
            "skills": "TypeScript · React · Styled Components · Immer",
        },
        {
            "title": "To-Do  —  Front-end",
            "bullets": [
                "Task manager with 9 standardized components published via Tsup.",
            ],
            "skills": "TypeScript · React",
        },
    ]


CV = {
    "pt": {
        "headline": ["Desenvolvedor Full Stack", "Gerente de Redes"],
        "contact_h": "CONTATO",
        "contact": [
            "julionev10@gmail.com",
            "+55 81 98627-3135",
            "linkedin.com/in/julionev10",
            "github.com/lionev",
            "Recife, PE",
        ],
        "dev_h": "HABILIDADES — DESENVOLVIMENTO",
        "dev": (
            "JavaScript · TypeScript · React · Next.js · Node.js · Vite · Tailwind CSS · "
            "Express · PostgreSQL · Prisma · JWT · LDAP · Python 3 · PySide6 · C# · .NET 8 · "
            "WinForms · WPF · OpenCV · PyInstaller · Cypress"
        ),
        "net_h": "HABILIDADES — REDES E SUPORTE",
        "net": (
            "TCP/IP · BGP · OSPF · MPLS · SD-WAN · Firewall · IDS/IPS · ZTNA · VPN · "
            "Active Directory · DNS · DHCP · Help Desk · Suporte N1/N2/N3 · ERP · "
            "Acesso remoto · Manutenção de hardware · Virtualização · Windows Server · "
            "Linux · ITIL · Gestão de infraestrutura · Liderança de equipe"
        ),
        "lang_h": "IDIOMAS",
        "langs": [("Português", "Nativo"), ("Inglês", "A2")],
        "edu_h": "EDUCAÇÃO",
        "edu": [
            ("Bacharelado em Sistemas da Informação", "UNINASSAU · 2024 · Concluído"),
            ("QA e Testes Automatizados", "Udemy · 2024 · Cypress básico e intermediário"),
            ("Desenvolvedor Front-End", "Rocketseat · 2024 · React, Java, NLWs"),
        ],
        "exp_h": "EXPERIÊNCIA PROFISSIONAL",
        "proj_h": "PROJETOS",
        "jobs": jobs_pt(),
        "projects": projects_pt(),
        "docx": "Julio_Jorge_Curriculo.docx",
        "pdf": "Julio_Jorge_Curriculo.pdf",
    },
    "en": {
        "headline": ["Full Stack Developer", "Network Manager"],
        "contact_h": "CONTACT",
        "contact": [
            "julionev10@gmail.com",
            "+55 81 98627-3135",
            "linkedin.com/in/julionev10",
            "github.com/lionev",
            "Recife, PE — Brazil",
        ],
        "dev_h": "SKILLS — DEVELOPMENT",
        "dev": (
            "JavaScript · TypeScript · React · Next.js · Node.js · Vite · Tailwind CSS · "
            "Express · PostgreSQL · Prisma · JWT · LDAP · Python 3 · PySide6 · C# · .NET 8 · "
            "WinForms · WPF · OpenCV · PyInstaller · Cypress"
        ),
        "net_h": "SKILLS — NETWORKING & SUPPORT",
        "net": (
            "TCP/IP · BGP · OSPF · MPLS · SD-WAN · Firewall · IDS/IPS · ZTNA · VPN · "
            "Active Directory · DNS · DHCP · Help Desk · Support N1/N2/N3 · ERP · "
            "Remote access · Hardware maintenance · Virtualization · Windows Server · "
            "Linux · ITIL · Infrastructure management · Team leadership"
        ),
        "lang_h": "LANGUAGES",
        "langs": [("Portuguese", "Native"), ("English", "A2")],
        "edu_h": "EDUCATION",
        "edu": [
            ("B.Sc. in Information Systems", "UNINASSAU · 2024 · Completed"),
            ("QA & automated testing", "Udemy · 2024 · Cypress beginner and intermediate"),
            ("Front-End Developer", "Rocketseat · 2024 · React, Java, NLWs"),
        ],
        "exp_h": "PROFESSIONAL EXPERIENCE",
        "proj_h": "PROJECTS",
        "jobs": jobs_en(),
        "projects": projects_en(),
        "docx": "Julio_Jorge_Resume.docx",
        "pdf": "Julio_Jorge_Resume.pdf",
    },
}


def rgb(hex_color: str) -> RGBColor:
    hex_color = hex_color.lstrip("#")
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def set_run_font(run, name="Calibri"):
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:cs"), name)
    rFonts.set(qn("w:eastAsia"), name)


def shade_cell(cell, fill: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_borders(cell, color="FFFFFF"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def set_cell_margins(cell, top, left, bottom, right):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for name, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{name}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        tcMar.append(el)
    tcPr.append(tcMar)


def set_cell_width(cell, width):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(width))
    tcW.set(qn("w:type"), "dxa")


def disable_cant_split(row):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "0")
    trPr.append(el)


def add_p(cell, text, *, size, color, bold=False, before=0, after=40, first=False):
    p = cell.paragraphs[0] if first and not cell.paragraphs[0].text else cell.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Twips(before)
    pf.space_after = Twips(after)
    pf.line_spacing = 1.08
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    set_run_font(run)
    return p


def build_docx(lang: str, dest: Path) -> Path:
    data = CV[lang]
    doc = Document()
    section = doc.sections[0]
    section.page_width = Twips(PAGE_W)
    section.page_height = Twips(PAGE_H)
    section.left_margin = Twips(0)
    section.right_margin = Twips(0)
    section.top_margin = Twips(0)
    section.bottom_margin = Twips(0)

    # Remove default empty paragraph
    body = doc.element.body
    for child in list(body):
        if child.tag == qn("w:p"):
            body.remove(child)

    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.allow_autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(PAGE_W))
    tblW.set(qn("w:type"), "dxa")

    grid = tbl.find(qn("w:tblGrid"))
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(1, grid)
    for child in list(grid):
        grid.remove(child)
    for w in (SIDEBAR_W, MAIN_W):
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(w))
        grid.append(col)

    left, right = table.rows[0].cells
    disable_cant_split(table.rows[0])

    for cell, width, fill, margins in (
        (left, SIDEBAR_W, "1A1A1A", (560, 420, 560, 380)),
        (right, MAIN_W, "FFFFFF", (480, 520, 480, 560)),
    ):
        set_cell_width(cell, width)
        shade_cell(cell, fill)
        set_cell_borders(cell)
        set_cell_margins(cell, *margins)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    add_p(left, "Julio Jorge", size=16, color="FFFFFF", bold=True, before=40, after=40, first=True)
    for i, line in enumerate(data["headline"]):
        add_p(left, line, size=8.5, color="AAAAAA", after=4 if i < len(data["headline"]) - 1 else 20)
    add_p(left, data["contact_h"], size=8, color="888888", bold=True, before=200, after=50)
    for i, line in enumerate(data["contact"]):
        add_p(left, line, size=9, color="BBBBBB", after=50)
    add_p(left, data["dev_h"], size=8, color="888888", bold=True, before=200, after=50)
    add_p(left, data["dev"], size=8, color="999999", after=50)
    add_p(left, data["net_h"], size=8, color="888888", bold=True, before=200, after=50)
    add_p(left, data["net"], size=8, color="999999", after=50)
    add_p(left, data["lang_h"], size=8, color="888888", bold=True, before=200, after=50)
    for name, level in data["langs"]:
        add_p(left, f"{name}  —  {level}", size=9, color="BBBBBB", after=40)
    add_p(left, data["edu_h"], size=8, color="888888", bold=True, before=180, after=50)
    for i, (title, meta) in enumerate(data["edu"]):
        add_p(left, title, size=9, color="BBBBBB", before=80 if i else 0, after=8)
        add_p(left, meta, size=8, color="666666", after=40)

    add_p(right, data["exp_h"], size=8.5, color="333333", bold=True, before=40, after=40, first=True)
    for job in data["jobs"]:
        add_p(right, job["title"], size=10, color="1A1A1A", bold=True, before=110, after=16)
        if job["subtitle"]:
            add_p(right, job["subtitle"], size=8, color="888888", after=16)
        for bullet in job["bullets"]:
            add_p(right, "•  " + bullet, size=9, color="555555", after=28)
        add_p(right, "Skills: " + job["skills"], size=8, color="888888", before=20, after=60)

    add_p(right, data["proj_h"], size=8.5, color="333333", bold=True, before=140, after=40)
    for proj in data["projects"]:
        add_p(right, proj["title"], size=10, color="1A1A1A", bold=True, before=90, after=16)
        for bullet in proj["bullets"]:
            add_p(right, "•  " + bullet, size=9, color="555555", after=20)
        add_p(right, "Skills: " + proj["skills"], size=8, color="888888", before=12, after=40)

    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(dest)
    return dest


def register_fonts():
    pdfmetrics.registerFont(TTFont("Calibri", str(FONTS / "calibri.ttf")))
    pdfmetrics.registerFont(TTFont("Calibri-Bold", str(FONTS / "calibrib.ttf")))


def wrap_canvas_text(c, text, x, y, max_width, font, size, leading, color):
    c.setFont(font, size)
    c.setFillColor(color)
    words = text.split()
    lines = []
    current = ""
    for word in words:
        trial = (current + " " + word).strip()
        if pdfmetrics.stringWidth(trial, font, size) <= max_width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    for line in lines:
        c.drawString(x, y, line)
        y -= leading
    return y


def draw_sidebar(c, data, page_width, page_height):
    sidebar_w = 2.15 * inch
    c.setFillColor(HexColor("#1A1A1A"))
    c.rect(0, 0, sidebar_w, page_height, fill=1, stroke=0)

    x = 0.32 * inch
    max_w = sidebar_w - 0.62 * inch
    y = page_height - 0.55 * inch
    white = HexColor("#FFFFFF")
    gray_a = HexColor("#AAAAAA")
    gray_8 = HexColor("#888888")
    gray_b = HexColor("#BBBBBB")
    gray_9 = HexColor("#999999")
    gray_6 = HexColor("#666666")

    c.setFillColor(white)
    c.setFont("Calibri-Bold", 16)
    c.drawString(x, y, "Julio Jorge")
    y -= 16
    for line in data["headline"]:
        y = wrap_canvas_text(c, line, x, y, max_w, "Calibri", 8.5, 11, gray_a)

    y -= 18
    c.setFillColor(gray_8)
    c.setFont("Calibri-Bold", 8)
    c.drawString(x, y, data["contact_h"])
    y -= 14
    for line in data["contact"]:
        y = wrap_canvas_text(c, line, x, y, max_w, "Calibri", 9, 12, gray_b)
        y -= 2

    y -= 12
    c.setFillColor(gray_8)
    c.setFont("Calibri-Bold", 8)
    c.drawString(x, y, data["dev_h"])
    y -= 13
    y = wrap_canvas_text(c, data["dev"], x, y, max_w, "Calibri", 8, 11, gray_9)

    y -= 14
    c.setFillColor(gray_8)
    c.setFont("Calibri-Bold", 8)
    c.drawString(x, y, data["net_h"])
    y -= 13
    y = wrap_canvas_text(c, data["net"], x, y, max_w, "Calibri", 8, 11, gray_9)

    y -= 14
    c.setFillColor(gray_8)
    c.setFont("Calibri-Bold", 8)
    c.drawString(x, y, data["lang_h"])
    y -= 13
    for name, level in data["langs"]:
        y = wrap_canvas_text(c, f"{name}  —  {level}", x, y, max_w, "Calibri", 9, 12, gray_b)
        y -= 2

    y -= 10
    c.setFillColor(gray_8)
    c.setFont("Calibri-Bold", 8)
    c.drawString(x, y, data["edu_h"])
    y -= 13
    for title, meta in data["edu"]:
        y = wrap_canvas_text(c, title, x, y, max_w, "Calibri", 9, 11, gray_b)
        y = wrap_canvas_text(c, meta, x, y, max_w, "Calibri", 8, 11, gray_6)
        y -= 6


def build_pdf(lang: str, dest: Path) -> Path:
    data = CV[lang]
    dest.parent.mkdir(parents=True, exist_ok=True)
    page_w, page_h = letter
    sidebar_w = 2.15 * inch

    styles = {
        "h": ParagraphStyle(
            "h",
            fontName="Calibri-Bold",
            fontSize=8.5,
            textColor=HexColor("#333333"),
            spaceBefore=2,
            spaceAfter=4,
            leading=11,
        ),
        "job": ParagraphStyle(
            "job",
            fontName="Calibri-Bold",
            fontSize=10,
            textColor=HexColor("#1A1A1A"),
            spaceBefore=9,
            spaceAfter=2,
            leading=13,
        ),
        "sub": ParagraphStyle(
            "sub",
            fontName="Calibri",
            fontSize=8,
            textColor=HexColor("#888888"),
            spaceBefore=0,
            spaceAfter=2,
            leading=11,
        ),
        "body": ParagraphStyle(
            "body",
            fontName="Calibri",
            fontSize=9,
            textColor=HexColor("#555555"),
            spaceBefore=0,
            spaceAfter=2.4,
            leading=12,
        ),
        "skills": ParagraphStyle(
            "skills",
            fontName="Calibri",
            fontSize=8,
            textColor=HexColor("#888888"),
            spaceBefore=1,
            spaceAfter=5,
            leading=11,
        ),
    }

    story = []
    story.append(Paragraph(data["exp_h"], styles["h"]))
    for job in data["jobs"]:
        block = [Paragraph(job["title"], styles["job"])]
        if job["subtitle"]:
            block.append(Paragraph(job["subtitle"], styles["sub"]))
        for bullet in job["bullets"]:
            block.append(Paragraph("•  " + bullet, styles["body"]))
        block.append(Paragraph("Skills: " + job["skills"], styles["skills"]))
        story.append(KeepTogether(block))

    story.append(Spacer(1, 8))
    story.append(Paragraph(data["proj_h"], styles["h"]))
    for proj in data["projects"]:
        block = [Paragraph(proj["title"], styles["job"])]
        for bullet in proj["bullets"]:
            block.append(Paragraph("•  " + bullet, styles["body"]))
        block.append(Paragraph("Skills: " + proj["skills"], styles["skills"]))
        story.append(KeepTogether(block))

    def on_page(canvas, doc):
        canvas.saveState()
        draw_sidebar(canvas, data, page_w, page_h)
        canvas.restoreState()

    right = Frame(
        sidebar_w,
        0.38 * inch,
        page_w - sidebar_w - 0.42 * inch,
        page_h - 0.76 * inch,
        leftPadding=16,
        rightPadding=8,
        topPadding=0,
        bottomPadding=0,
        id="main",
    )
    doc = BaseDocTemplate(
        str(dest),
        pagesize=letter,
        leftMargin=0,
        rightMargin=0,
        topMargin=0,
        bottomMargin=0,
        title="Julio Jorge — Curriculum Vitae" if lang == "pt" else "Julio Jorge — Resume",
        author="Julio Jorge",
    )
    doc.addPageTemplates([PageTemplate(id="cv", frames=[right], onPage=on_page)])
    doc.build(story)
    return dest


def main():
    register_fonts()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    outputs = []
    for lang in ("pt", "en"):
        data = CV[lang]
        docx_path = build_docx(lang, OUT_DIR / data["docx"])
        pdf_path = build_pdf(lang, OUT_DIR / data["pdf"])
        outputs.extend([docx_path, pdf_path])

    PUBLIC.mkdir(parents=True, exist_ok=True)
    shutil.copy2(OUT_DIR / CV["pt"]["pdf"], PUBLIC / "Julio_Jorge.pdf")
    shutil.copy2(OUT_DIR / CV["en"]["pdf"], PUBLIC / "Julio_Jorge_Resume.pdf")

    for path in outputs:
        shutil.copy2(path, DOWNLOADS / path.name)

    print("Generated:")
    for path in outputs:
        print(" ", path, path.stat().st_size, "bytes")
        print("  copy:", DOWNLOADS / path.name)


if __name__ == "__main__":
    main()
